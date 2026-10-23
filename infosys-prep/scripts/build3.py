import csv, json, glob, re
from collections import defaultdict, Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

D = "/mnt/project-files/infosys-prep/"
inf = list(csv.DictReader(open(D + "infosys_sp_dse_questions.csv", encoding="utf-8")))
cf = list(csv.DictReader(open(D + "cf/cf_dp_plus_greedy.csv", encoding="utf-8")))

def load(pattern, key):
    out = {}
    for f in sorted(glob.glob(D + "raw_notes/" + pattern)):
        for line in open(f, encoding="utf-8", errors="replace"):
            try:
                o = json.loads(line)
            except Exception:
                continue
            if o.get(key):
                out[o[key]] = o
    return out

inotes, cnotes = load("inf_notes*.jsonl", "id"), load("cf_notes*.jsonl", "cf_id")

def pats(n):
    p = n.get("patterns") or []
    return [p] if isinstance(p, str) else [x for x in p if isinstance(x, str)]

def num(x, d=None):
    try:
        return int(float(x))
    except Exception:
        return d

# ---- Infosys question weight: how much a question says about what the exam asks ----
CLAIM = {"actual": 1.0, "memory-based recollection": 0.7, "previous-year compilation": 0.4}
def recency(years):
    ys = [num(y) for y in years.split(",") if num(y)]
    if not ys: return 0.5
    y = max(ys)
    return 1.0 if y >= 2025 else 0.8 if y >= 2023 else 0.6
def round_w(r):
    r = (r or "").lower()
    return 0.6 if "interview" in r else 0.8 if r in ("", "unknown") else 1.0
for q in inf:
    n = inotes.get(q["id"], {})
    q["_p"] = pats(n)
    ns = num(q["n_sources"], 1) or 1
    q["_w"] = CLAIM.get(q["source_claim"], 0.5) * recency(q["years"]) * round_w(q["round"]) * min(2.0, 1 + 0.25 * (ns - 1))

R = defaultdict(float)
for q in inf:
    for p in q["_p"]:
        R[p] += q["_w"] / len(q["_p"])
Rmax = max(R.values()) if R else 1
Rn = {p: v / Rmax for p, v in R.items()}
wmax = max(q["_w"] for q in inf)
for q in inf:
    rp = max([Rn.get(p, 0) for p in q["_p"]] or [0])
    q["_pred"] = round(100 * (0.7 * rp + 0.3 * q["_w"] / wmax))

# recurrence among harder (CF 1300+) Infosys questions, used to score the CF list
R2 = defaultdict(float)
for q in inf:
    if (num(q["cf_rating"]) or 0) >= 1300:
        for p in q["_p"]:
            R2[p] += q["_w"] / len(q["_p"])
R2max = max(R2.values()) if R2 else 1
R2n = {p: v / R2max for p, v in R2.items()}

# rating fit: how many Infosys OA-style questions sit near this rating
oa = [num(q["cf_rating"]) for q in inf if num(q["cf_rating"]) and "interview" not in q["round"].lower()]
def fit_raw(r): return sum(1 for x in oa if abs(x - r) <= 150)
fmax = max(fit_raw(r) for r in range(800, 3001, 100)) or 1

TAGMAP = {"dp": "DP", "greedy": "Greedy", "graphs": "Graphs", "dfs and similar": "Graphs", "shortest paths": "Graphs",
          "trees": "Trees", "strings": "Strings", "math": "Math/Number Theory", "number theory": "Math/Number Theory",
          "binary search": "Binary Search", "bitmasks": "Bit Manipulation", "two pointers": "Two Pointers/Sliding Window",
          "sortings": "Sorting", "combinatorics": "Combinatorics", "dsu": "DSU", "implementation": "Implementation",
          "brute force": "Implementation", "hashing": "Hashing", "data structures": "Segment Tree/BIT", "matrices": "Matrix"}
def jac(a, b):
    a, b = set(a), set(b)
    return len(a & b) / len(a | b) if a and b else 0.0

infp = [(q, set(q["_p"]), {t.strip() for t in q["all_topics"].split(",") if t.strip()}, num(q["cf_rating"])) for q in inf if (num(q["cf_rating"]) or 0) >= 1200]
for c in cf:
    n = cnotes.get(c["cf_id"], {})
    c["_n"] = n
    cp = set(pats(n))
    ct = {TAGMAP[t.strip()] for t in c["tags"].split(",") if t.strip() in TAGMAP}
    r = num(c["rating"], 1500)
    best, bq = 0.0, None
    for q, qp, qt, qr in infp:
        if not qp and not qt: continue
        rc = max(0.0, 1 - abs(r - qr) / 800) if qr else 0.5
        s = 0.6 * jac(cp, qp) + 0.25 * jac(ct, qt) + 0.15 * rc
        if s > best: best, bq = s, q
    c["_match"] = round(100 * best)
    c["_near"] = f'{bq["id"]} {bq["title"]}' if bq else ""
    rp = max([R2n.get(p, 0) for p in cp] or [0])
    c["_pred"] = round(100 * (0.5 * best + 0.35 * rp + 0.15 * fit_raw(r) / fmax))

def j(v):
    return "\n".join(v) if isinstance(v, list) else (v or "")

# ---- outputs ----
icols = ["id", "title", "primary_topic", "patterns", "cf_rating", "prediction_index", "years", "roles", "round", "source_claim",
         "brief", "statement", "constraints", "sample", "key_idea", "solution", "complexity", "dry_run", "cpp", "confidence",
         "n_sources", "source_urls", "notes"]
irows = []
for q in inf:
    n = inotes.get(q["id"], {})
    irows.append({**{k: q.get(k, "") for k in icols}, "patterns": ", ".join(q["_p"]), "prediction_index": q["_pred"],
                  "brief": n.get("brief", ""), "key_idea": n.get("key_idea", ""), "solution": j(n.get("solution")),
                  "complexity": n.get("complexity", ""), "dry_run": j(n.get("dry_run")), "cpp": n.get("cpp", ""),
                  "confidence": n.get("confidence", "")})
ccols = ["cf_id", "name", "rating", "tags", "patterns", "match_pct", "nearest_infosys_question", "prediction_index",
         "brief", "key_idea", "solution", "complexity", "dry_run", "cpp", "confidence", "solved", "url"]
crows = []
for c in cf:
    n = c["_n"]
    crows.append({"cf_id": c["cf_id"], "name": c["name"], "rating": num(c["rating"]), "tags": c["tags"], "patterns": ", ".join(pats(n)),
                  "match_pct": c["_match"], "nearest_infosys_question": c["_near"], "prediction_index": c["_pred"],
                  "brief": n.get("brief", ""), "key_idea": n.get("key_idea", ""), "solution": j(n.get("solution")),
                  "complexity": n.get("complexity", ""), "dry_run": j(n.get("dry_run")), "cpp": n.get("cpp", ""),
                  "confidence": n.get("confidence", ""), "solved": num(c["solved"]), "url": c["url"]})
crows.sort(key=lambda r: -r["prediction_index"])

for name, cols, rows in [("infosys_sp_dse_questions_full.csv", icols, irows), ("cf_dp_scored.csv", ccols, crows)]:
    with open(D + name, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=cols); w.writeheader(); w.writerows(rows)

wb = Workbook()
hdr, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
WID = {"title": 30, "name": 28, "brief": 45, "statement": 50, "solution": 55, "dry_run": 45, "cpp": 60, "key_idea": 35,
       "nearest_infosys_question": 32, "source_urls": 40, "url": 30, "notes": 30, "patterns": 24}
def sheet(title, cols, rows, link=None):
    ws = wb.create_sheet(title[:31])
    ws.append(cols)
    for cell in ws[1]: cell.font, cell.fill = hdr, fill
    for r in rows:
        ws.append([r[c] for c in cols])
        if link:
            cell = ws.cell(row=ws.max_row, column=cols.index(link) + 1)
            u = str(r[link]).split(" | ")[0]
            if u.startswith("http"): cell.hyperlink, cell.font = u, Font(color="0563C1", underline="single")
    for i, c in enumerate(cols, 1): ws.column_dimensions[get_column_letter(i)].width = WID.get(c, 12)
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions

s = wb.active; s.title = "Summary"
s.append(["Infosys SP/DSE prep database"]); s["A1"].font = Font(bold=True, size=14)
s.append([f"Infosys questions: {len(irows)} (notes written for {sum(1 for r in irows if r['solution'])})"])
s.append([f"Codeforces DP problems from your list: {len(crows)} (notes written for {sum(1 for r in crows if r['solution'])})"])
s.append([])
s.append(["Pattern", "Infosys recurrence score (0-100)", "Infosys questions with this pattern", "CF problems with this pattern"])
for cell in s[5]: cell.font, cell.fill = hdr, fill
ic = Counter(p for q in inf for p in q["_p"]); cc = Counter(p for c in cf for p in pats(c["_n"]))
for p, v in sorted(Rn.items(), key=lambda x: -x[1]):
    s.append([p, round(100 * v), ic[p], cc[p]])
s.column_dimensions["A"].width = 40

m = wb.create_sheet("How scores work")
for line in [
    "These scores are heuristics built from the collected Infosys data. They rank what to practise; they are not true probabilities.",
    "",
    "Question weight w (Infosys questions): source claim (actual 1.0, recollection 0.7, prep-site compilation 0.4) x recency (2025-26: 1.0, 2023-24: 0.8, 2021-22: 0.6, unknown 0.5) x round (OA/HackWithInfy 1.0, unknown 0.8, interview 0.6) x repeats (+25% per extra source, max x2).",
    "Pattern recurrence R(p): sum of w over Infosys questions with pattern p, scaled so the most frequent pattern = 100.",
    "Infosys prediction_index = 0.7 x (highest R of its patterns) + 0.3 x (its own weight / max weight), on 0-100.",
    "",
    "CF scores only compare against Infosys questions rated CF 1200+, and use pattern recurrence among Infosys questions rated 1300+, since your CF list is 1400-2200.",
    "CF match_pct = best similarity to any such Infosys question, where similarity = 0.6 x pattern overlap (Jaccard) + 0.25 x topic-tag overlap + 0.15 x rating closeness (1 - |rating gap|/800). nearest_infosys_question names that question.",
    "CF prediction_index = 0.5 x match + 0.35 x (highest R of its patterns) + 0.15 x rating fit (how many Infosys OA questions sit within +-150 of its rating), on 0-100.",
    "",
    "Solutions, dry runs and C++ were written by Claude from the source statement (Infosys) or the Codeforces statement page (CF). Check the confidence column; 'low' means the source statement was incomplete.",
]:
    m.append([line])
m.column_dimensions["A"].width = 160
for row in m.iter_rows():
    for cell in row: cell.alignment = Alignment(wrap_text=True)

sheet("Infosys Questions", icols, sorted(irows, key=lambda r: -r["prediction_index"]), "source_urls")
for t in [t for t, _ in Counter(r["primary_topic"] for r in irows).most_common()]:
    sheet("INF " + t.replace("/", "-"), icols, sorted([r for r in irows if r["primary_topic"] == t], key=lambda r: -r["prediction_index"]), "source_urls")
sheet("CF DP Problems", ccols, crows, "url")
sheet("CF Top 150 to practise", ccols, crows[:150], "url")
wb.save(D + "infosys_sp_dse_prep_full.xlsx")
print("inf notes", len(inotes), "cf notes", len(cnotes), "patterns", len(R))
